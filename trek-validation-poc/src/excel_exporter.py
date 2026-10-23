"""CFO-ready ROI workbook. Every financial output cell is a live Excel formula."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

TEAL = "0F766E"
INK = "10161D"
LIGHT_TEAL = "E6F4F1"
INPUT_FILL = "FFF7E0"
CURRENCY = '"$"#,##0'
CURRENCY_CENTS = '"$"#,##0.00'
PCT = "0.0%"
MULTIPLE = '0.0"x"'
MONTHS = '0.0" mo"'

SHEET_EXEC = "Executive_ROI_Business_Case"
SHEET_BENCH = "Fee_Schedule_Benchmark"
SHEET_DNF = "Data_Not_Found_Resolution"

THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=TEAL)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT_TEAL)

BENCH_HEADERS = [
    "Row ID", "Payer", "CPT", "Description", "Service Line", "Prospect Rate", "Medicare Rate",
    "Market P50", "Market P75", "Current Annual Revenue ($)", "Annual Uplift to P50 ($)",
    "Annual Uplift to P75 ($)", "Annual Volume", "Variance Flag", "Prospect % of Medicare",
    "Market P50 % of Medicare", "Unit Rate Gap ($)", "Peer MRF Citation",
]
DNF_HEADERS = [
    "Row ID", "Payer", "Unmatched CPT", "Modifier", "Prospect Rate", "Current Baseline Rev ($)",
    "Annual Volume", "DNF Reason Code", "Proxy CPT", "Proxy Market P50", "Annual DNF Recovery ($)",
    "Adjustment Factor", "Estimated Benchmark Rate", "Confidence Tier", "Proxy MRF Citation",
]


def _header_row(ws, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def _widths(ws, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _build_benchmark(ws, v: pd.DataFrame) -> int:
    _header_row(ws, BENCH_HEADERS)
    n = len(v)
    for i, (_, r) in enumerate(v.iterrows()):
        row = i + 2
        values = {
            "A": r["prospect_row_id"], "B": r["payer_name"], "C": str(r["cpt_code"]), "D": r["description"],
            "E": r["service_line"], "F": round(float(r["prospect_rate"]), 2), "G": round(float(r["medicare_rate"]), 2),
            "H": float(r["market_p50"]), "I": float(r["market_p75"]),
            "J": f"=F{row}*M{row}",
            "K": f"=IF(H{row}>F{row},(H{row}-F{row})*M{row},0)",
            "L": f"=IF(I{row}>F{row},(I{row}-F{row})*M{row},0)",
            "M": int(r["annual_volume"]), "N": r["variance_flag"],
            "O": f"=IF(G{row}>0,F{row}/G{row},0)",
            "P": f"=IF(G{row}>0,H{row}/G{row},0)",
            "Q": f"=H{row}-F{row}",
            "R": r.get("sample_mrf_citation", ""),
        }
        for col, val in values.items():
            cell = ws[f"{col}{row}"]
            cell.value = val
            cell.border = BORDER
        for col in ("F", "G", "H", "I", "Q"):
            ws[f"{col}{row}"].number_format = CURRENCY_CENTS
        for col in ("J", "K", "L"):
            ws[f"{col}{row}"].number_format = CURRENCY
        ws[f"M{row}"].number_format = "#,##0"
        for col in ("O", "P"):
            ws[f"{col}{row}"].number_format = PCT
    last = n + 1
    total_row = n + 2
    ws[f"A{total_row}"] = "TOTAL / PORTFOLIO AVERAGE"
    totals = {
        "J": f"=SUM(J2:J{last})", "K": f"=SUM(K2:K{last})", "L": f"=SUM(L2:L{last})",
        "M": f"=SUM(M2:M{last})", "O": f"=AVERAGE(O2:O{last})", "P": f"=AVERAGE(P2:P{last})",
    }
    for col, f in totals.items():
        ws[f"{col}{total_row}"] = f
    for c in range(1, len(BENCH_HEADERS) + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(bold=True, color=INK)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    for col in ("J", "K", "L"):
        ws[f"{col}{total_row}"].number_format = CURRENCY
    ws[f"M{total_row}"].number_format = "#,##0"
    for col in ("O", "P"):
        ws[f"{col}{total_row}"].number_format = PCT
    _widths(ws, {"A": 11, "B": 22, "C": 8, "D": 36, "E": 20, "F": 13, "G": 13, "H": 13, "I": 13, "J": 18,
                 "K": 18, "L": 18, "M": 12, "N": 25, "O": 13, "P": 13, "Q": 14, "R": 14})
    ws.auto_filter.ref = f"A1:R{last}"
    return last


def _build_dnf(ws, d: pd.DataFrame) -> int:
    _header_row(ws, DNF_HEADERS)
    n = len(d)
    for i, (_, r) in enumerate(d.iterrows()):
        row = i + 2
        resolved = pd.notna(r["estimated_market_rate"])
        values = {
            "A": r["prospect_row_id"], "B": r["payer_name"], "C": str(r["cpt_code"]), "D": r["modifier"],
            "E": round(float(r["prospect_rate"]), 2),
            "F": f"=E{row}*G{row}",
            "G": int(r["annual_volume"]), "H": r["dnf_reason_code"], "I": r["proxy_cpt"],
            "J": float(r["proxy_market_p50"]) if resolved else 0,
            "K": f"=IF(M{row}>E{row},(M{row}-E{row})*G{row},0)",
            "L": float(r["adjustment_factor"]) if resolved else 0,
            "M": f"=J{row}*L{row}",
            "N": r["confidence_tier"], "O": r["proxy_mrf_citation"],
        }
        for col, val in values.items():
            cell = ws[f"{col}{row}"]
            cell.value = val
            cell.border = BORDER
        for col in ("E", "J", "M"):
            ws[f"{col}{row}"].number_format = CURRENCY_CENTS
        for col in ("F", "K"):
            ws[f"{col}{row}"].number_format = CURRENCY
        ws[f"L{row}"].number_format = '0.00"x"'
        ws[f"L{row}"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    last = n + 1
    total_row = n + 2
    ws[f"A{total_row}"] = "TOTAL"
    ws[f"F{total_row}"] = f"=SUM(F2:F{last})"
    ws[f"K{total_row}"] = f"=SUM(K2:K{last})"
    for c in range(1, len(DNF_HEADERS) + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(bold=True, color=INK)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    for col in ("F", "K"):
        ws[f"{col}{total_row}"].number_format = CURRENCY
    _widths(ws, {"A": 11, "B": 22, "C": 13, "D": 10, "E": 13, "F": 18, "G": 12, "H": 32, "I": 10, "J": 15,
                 "K": 20, "L": 13, "M": 17, "N": 28, "O": 16})
    return last


def _build_exec(ws, bench_last: int, dnf_last: int, params: dict) -> None:
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Trek Health — Payer Performance & Validation Business Case"
    ws["B2"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["B2"].fill = HEADER_FILL
    ws.merge_cells("B2:D2")
    ws["B3"] = params.get("prospect_name", "Summit Regional Health System (Synthetic Validation File)")
    ws["B3"].font = Font(italic=True, color=INK)

    ws["B5"] = "A · Editable Executive Assumptions (yellow cells drive every number below)"
    ws["B5"].font = Font(bold=True, color=TEAL, size=12)
    assumptions = [
        (6, "Annual Trek Platform Investment ($)", float(params.get("annual_subscription_cost", 85000)), CURRENCY),
        (7, "Negotiation Realization Factor", float(params.get("realization_factor", 0.5)), PCT),
        (8, "Include DNF Crosswalk Recovery (1 = Yes, 0 = No)", 1 if params.get("include_dnf", True) else 0, "0"),
        (9, "Contract Term (Years)", int(params.get("term_years", 3)), "0"),
    ]
    for row, label, val, fmt in assumptions:
        ws[f"B{row}"] = label
        ws[f"C{row}"] = val
        ws[f"C{row}"].number_format = fmt
        ws[f"C{row}"].fill = PatternFill("solid", fgColor=INPUT_FILL)
        for col in ("B", "C"):
            ws[f"{col}{row}"].border = BORDER

    ws["B11"] = "B · Live-Formula Financial Summary"
    ws["B11"].font = Font(bold=True, color=TEAL, size=12)
    b, d = SHEET_BENCH, SHEET_DNF
    summary = [
        (12, "Current Contracted Baseline Revenue", f"=SUM({b}!J2:J{bench_last})+SUM({d}!F2:F{dnf_last})", CURRENCY),
        (13, "Validated Gross Uplift at Market P50", f"=SUM({b}!K2:K{bench_last})", CURRENCY),
        (14, "Validated Gross Uplift at Market P75", f"=SUM({b}!L2:L{bench_last})", CURRENCY),
        (15, "DNF Quantified Recovery Opportunity", f"=SUM({d}!K2:K{dnf_last})*C8", CURRENCY),
        (16, "Total Addressable Annual Opportunity (P50 + DNF)", "=C13+C15", CURRENCY),
        (17, "Net Realized Annual Revenue Uplift", "=C16*C7", CURRENCY),
        (18, "Year 1 Net Financial Benefit", "=C17-C6", CURRENCY),
        (19, "First-Year ROI Multiple", "=IF(C6>0,C17/C6,0)", MULTIPLE),
        (20, "Payback Period (Months)", "=IF(C17>0,(C6/C17)*12,0)", MONTHS),
        (21, "Multi-Year Cumulative Net Value", "=(C17*C9)-(C6*C9)", CURRENCY),
    ]
    for row, label, formula, fmt in summary:
        ws[f"B{row}"] = label
        ws[f"C{row}"] = formula
        ws[f"C{row}"].number_format = fmt
        for col in ("B", "C"):
            ws[f"{col}{row}"].border = BORDER
        if row in (16, 17, 19, 20):
            ws[f"B{row}"].font = Font(bold=True, color=INK)
            ws[f"C{row}"].font = Font(bold=True, color=TEAL)
    ws["D12"] = "Sum of prospect rate × volume (validated + DNF lines)"
    ws["D13"] = "Σ max(0, P50 − prospect) × volume"
    ws["D15"] = "Σ max(0, proxy P50 × factor − prospect) × volume; toggled by C8"
    ws["D17"] = "Conservative: only the realization share of the gap is assumed won"
    ws["D20"] = "Months of realized uplift needed to cover the platform fee"
    for r in range(12, 22):
        ws[f"D{r}"].font = Font(italic=True, color="64748B", size=9)
    ws["B23"] = ("Methodology: Market percentiles computed from a normalized Transparency in Coverage MRF slice "
                 "(synthetic, 2026-Q3). DNF lines resolved by documented crosswalks — see Data_Not_Found_Resolution.")
    ws["B23"].font = Font(italic=True, color="64748B", size=9)
    ws.merge_cells("B23:D24")
    ws["B23"].alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 3, "B": 50, "C": 20, "D": 62})


def generate_roi_workbook(validated_df: pd.DataFrame, dnf_df: pd.DataFrame, scenario_params: dict | None = None) -> bytes:
    """Build the 3-tab ROI workbook and return .xlsx bytes."""
    params = scenario_params or {}
    wb = Workbook()
    ws_exec = wb.active
    ws_exec.title = SHEET_EXEC
    ws_bench = wb.create_sheet(SHEET_BENCH)
    ws_dnf = wb.create_sheet(SHEET_DNF)
    bench_last = _build_benchmark(ws_bench, validated_df)
    dnf_last = _build_dnf(ws_dnf, dnf_df)
    _build_exec(ws_exec, bench_last, dnf_last, params)
    ws_exec.sheet_properties.tabColor = TEAL
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def formula_preview(xlsx_bytes: bytes) -> pd.DataFrame:
    """Key formulas pulled back out of the generated workbook (for on-screen transparency)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    rows = []
    ex = wb[SHEET_EXEC]
    for r in range(12, 22):
        rows.append({"Sheet": SHEET_EXEC, "Cell": f"C{r}", "Label": ex[f"B{r}"].value, "Formula": ex[f"C{r}"].value})
    bench = wb[SHEET_BENCH]
    total_row = bench.max_row
    for col in ("J", "K", "L", "O", "P", "Q"):
        rows.append({"Sheet": SHEET_BENCH, "Cell": f"{col}2", "Label": bench[f"{col}1"].value, "Formula": bench[f"{col}2"].value})
    for col in ("J", "K", "L"):
        rows.append({"Sheet": SHEET_BENCH, "Cell": f"{col}{total_row}", "Label": f"TOTAL {bench[f'{col}1'].value}",
                     "Formula": bench[f"{col}{total_row}"].value})
    dnf = wb[SHEET_DNF]
    dnf_total = dnf.max_row
    for col in ("F", "K", "M"):
        rows.append({"Sheet": SHEET_DNF, "Cell": f"{col}2", "Label": dnf[f"{col}1"].value, "Formula": dnf[f"{col}2"].value})
    rows.append({"Sheet": SHEET_DNF, "Cell": f"K{dnf_total}", "Label": "TOTAL Annual DNF Recovery ($)",
                 "Formula": dnf[f"K{dnf_total}"].value})
    return pd.DataFrame(rows)


def count_formulas(xlsx_bytes: bytes) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    return sum(
        1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )

