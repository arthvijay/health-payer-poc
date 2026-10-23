"""Independent evaluator for the workbook's formulas.

openpyxl writes formulas but never computes them. This evaluator recomputes the
exact formula subset the exporter emits (arithmetic, comparisons, SUM, AVERAGE,
IF, same-sheet and cross-sheet references) so the UI can show evaluated values
and the test suite can prove the workbook math equals the engine math.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

_TOKEN = re.compile(
    r"\s*(?:"
    r"(?P<num>\d+(?:\.\d+)?)"
    r"|(?P<ref>(?:[A-Za-z_][A-Za-z0-9_]*!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"
    r"|(?P<func>[A-Z]+)\("
    r"|(?P<op><=|>=|<>|[-+*/(),<>=])"
    r")"
)


class WorkbookEvaluator:
    def __init__(self, xlsx_bytes: bytes):
        self.wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
        self.cache: dict[tuple[str, str], object] = {}

    def value(self, sheet: str, cell: str):
        key = (sheet, cell.replace("$", ""))
        if key in self.cache:
            return self.cache[key]
        raw = self.wb[sheet][key[1]].value
        if isinstance(raw, str) and raw.startswith("="):
            result = _Parser(self, sheet, raw[1:]).parse()
        else:
            result = raw
        self.cache[key] = result
        return result

    def _range(self, sheet: str, a: str, b: str) -> list:
        ca, ra = _split(a)
        cb, rb = _split(b)
        out = []
        for col in range(column_index_from_string(ca), column_index_from_string(cb) + 1):
            for row in range(ra, rb + 1):
                out.append(self.value(sheet, f"{get_column_letter(col)}{row}"))
        return out


def _split(ref: str) -> tuple[str, int]:
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", ref)
    return m.group(1), int(m.group(2))


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


class _Parser:
    def __init__(self, ev: WorkbookEvaluator, sheet: str, text: str):
        self.ev, self.sheet = ev, sheet
        self.tokens = self._tokenize(text)
        self.i = 0

    @staticmethod
    def _tokenize(text: str) -> list[tuple[str, str]]:
        pos, toks = 0, []
        text = text.strip()
        while pos < len(text):
            m = _TOKEN.match(text, pos)
            if not m or m.end() == pos:
                raise ValueError(f"Unsupported formula syntax near: {text[pos:]!r}")
            kind = m.lastgroup
            toks.append((kind, m.group(kind)))
            pos = m.end()
        return toks

    def peek(self):
        return self.tokens[self.i] if self.i < len(self.tokens) else (None, None)

    def take(self, expected: str | None = None):
        tok = self.peek()
        if expected is not None and tok[1] != expected:
            raise ValueError(f"Expected {expected!r}, got {tok[1]!r}")
        self.i += 1
        return tok

    def parse(self):
        v = self.comparison()
        if self.i != len(self.tokens):
            raise ValueError(f"Trailing tokens: {self.tokens[self.i:]}")
        return v

    def comparison(self):
        left = self.additive()
        while self.peek()[1] in ("<", ">", "<=", ">=", "=", "<>"):
            op = self.take()[1]
            right = self.additive()
            a, b = _num(left), _num(right)
            left = {"<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b, "=": a == b, "<>": a != b}[op]
        return left

    def additive(self):
        left = self.term()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            right = self.term()
            left = _num(left) + _num(right) if op == "+" else _num(left) - _num(right)
        return left

    def term(self):
        left = self.unary()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            right = self.unary()
            left = _num(left) * _num(right) if op == "*" else (_num(left) / _num(right) if _num(right) else 0.0)
        return left

    def unary(self):
        if self.peek()[1] == "-":
            self.take()
            return -_num(self.unary())
        return self.atom()

    def atom(self):
        kind, val = self.take()
        if kind == "num":
            return float(val)
        if kind == "ref":
            sheet, ref = (val.split("!", 1) if "!" in val else (self.sheet, val))
            if ":" in ref:
                a, b = ref.split(":")
                return self.ev._range(sheet, a.replace("$", ""), b.replace("$", ""))
            return self.ev.value(sheet, ref)
        if kind == "func":
            args = []
            if self.peek()[1] != ")":
                args.append(self.comparison())
                while self.peek()[1] == ",":
                    self.take()
                    args.append(self.comparison())
            self.take(")")
            return self._call(val, args)
        if val == "(":
            v = self.comparison()
            self.take(")")
            return v
        raise ValueError(f"Unexpected token {val!r}")

    @staticmethod
    def _flatten(args) -> list[float]:
        out = []
        for a in args:
            if isinstance(a, list):
                out.extend(_num(x) for x in a if x is not None and not isinstance(x, str))
            else:
                out.append(_num(a))
        return out

    def _call(self, name: str, args: list):
        if name == "SUM":
            return sum(self._flatten(args))
        if name == "AVERAGE":
            vals = self._flatten(args)
            return sum(vals) / len(vals) if vals else 0.0
        if name == "IF":
            cond = args[0]
            return args[1] if bool(cond) else (args[2] if len(args) > 2 else False)
        raise ValueError(f"Unsupported function {name}")
