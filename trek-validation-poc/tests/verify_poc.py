"""End-to-end verification suite for the Trek Health validation POC.

Run either way:
    pytest -v tests/verify_poc.py
    python tests/verify_poc.py
"""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src import agent as agent_mod  # noqa: E402
from src.agent import CFO_HEADER, PRESET_PROMPTS, TECH_HEADER, ask_benchmark_agent, load_skill  # noqa: E402
from src.data_generator import MRF_PATH, POLICIES_PATH, PROSPECT_PATH, generate_all  # noqa: E402
from src.demo_script import SCRIPT_PATH, SLIDES, default_facts, render_markdown  # noqa: E402
from src.engine import (CROSSWALK_LIBRARY, TARGET_FLOOR, TARGET_P75, TrekValidationEngine, compute_kpis,  # noqa: E402
                        load_prospect_csv, pareto_table, sensitivity_matrix, variance_flag)
from src.excel_exporter import SHEET_BENCH, SHEET_DNF, SHEET_EXEC, count_formulas, generate_roi_workbook  # noqa: E402
from src.formula_audit import WorkbookEvaluator  # noqa: E402

MRF_COLUMNS = ["mrf_record_id", "hospital_name", "ein_npi", "payer_name", "plan_type", "cpt_code", "description",
               "service_line", "billing_class", "medicare_rate", "negotiated_rate", "pct_of_medicare",
               "mrf_schema_version", "last_updated"]
PROSPECT_COLUMNS = ["prospect_row_id", "payer_name", "cpt_code", "modifier", "description", "service_line",
                    "prospect_rate", "annual_volume", "contract_renewal_date"]
EXPECTED_DNF = {
    "27446": ("CODE_ABSENT_FROM_PAYER_MRF", "27447", 0.82),
    "70551": ("MODIFIER_26_SPLIT_MISSING", "70553", 0.38),
    "45385": ("LOW_VOLUME_SUPPRESSED_IN_TIC", "45380", 1.14),
    "0408T": ("EMERGING_CAT_III_NO_MRF_RATE", "93458", 1.20),
}
SUB = 85_000.0


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def engine() -> TrekValidationEngine:
    generate_all()
    return TrekValidationEngine()


@pytest.fixture(scope="session")
def result(engine):
    return engine.run_validation(engine.default_prospect())


@pytest.fixture(scope="session")
def kpis(result):
    return compute_kpis(result.validated_df, result.dnf_df, realization_rate=0.5, annual_trek_subscription=SUB,
                        total_lines=result.total_lines)


@pytest.fixture(scope="session")
def xlsx(result):
    return generate_roi_workbook(result.validated_df, result.dnf_df,
                                 {"annual_subscription_cost": SUB, "realization_factor": 0.5, "include_dnf": True,
                                  "term_years": 3})


# --------------------------------------------------------------------------- 1. data generation
def test_data_files_exist_with_exact_schemas():
    mrf = pd.read_csv(MRF_PATH, dtype={"cpt_code": str})
    pro = pd.read_csv(PROSPECT_PATH, dtype={"cpt_code": str, "modifier": str}, keep_default_na=False)
    assert list(mrf.columns) == MRF_COLUMNS
    assert list(pro.columns) == PROSPECT_COLUMNS
    assert len(mrf) >= 500
    assert len(pro) == 60
    assert pro["cpt_code"].nunique() == 30
    assert set(pro["payer_name"]) == {"UnitedHealthcare (UHC)", "Aetna"}
    assert pro.groupby("cpt_code").size().eq(2).all()
    assert mrf["mrf_record_id"].is_unique and pro["prospect_row_id"].is_unique
    assert mrf["hospital_name"].nunique() == 5 and mrf["payer_name"].nunique() == 4 and mrf["cpt_code"].nunique() == 26


def test_market_rates_realistic_band():
    mrf = pd.read_csv(MRF_PATH, dtype={"cpt_code": str})
    ratio = mrf["negotiated_rate"] / mrf["medicare_rate"] * 100
    assert ratio.min() >= 144.99 and ratio.max() <= 330.01


def test_policies_json_valid():
    pol = json.loads(POLICIES_PATH.read_text())
    for cpt in ("27447", "70553", "45380", "93458", "29827", "27446", "45385"):
        for payer in ("UnitedHealthcare (UHC)", "Aetna"):
            card = pol[cpt][payer]
            assert {"prior_auth_required", "policy_change_2026", "common_denial_code", "denial_risk"} <= set(card)
            assert card["denial_risk"] in {"High", "Medium", "Low"}


def test_generator_is_deterministic(tmp_path):
    paths = generate_all(tmp_path)
    assert paths["mrf"].read_bytes() == MRF_PATH.read_bytes()
    assert paths["prospect"].read_bytes() == PROSPECT_PATH.read_bytes()
    assert paths["policies"].read_bytes() == POLICIES_PATH.read_bytes()


# --------------------------------------------------------------------------- 2. engine
def test_partition_52_validated_8_dnf_zero_dropped(result):
    v, d = result.validated_df, result.dnf_df
    assert len(v) == 52 and len(d) == 8
    assert len(v) + len(d) == result.total_lines == 60
    ids = set(v["prospect_row_id"]) | set(d["prospect_row_id"])
    assert ids == set(pd.read_csv(PROSPECT_PATH)["prospect_row_id"])
    assert set(v["prospect_row_id"]).isdisjoint(set(d["prospect_row_id"]))


def test_duckdb_percentiles_match_independent_pandas(engine):
    mrf = engine.mrf_df
    for _, c in engine.cohort_df.iterrows():
        rates = mrf[(mrf["cpt_code"] == c["cpt_code"]) & (mrf["payer_name"] == c["payer_name"])]["negotiated_rate"]
        assert c["market_p50"] == pytest.approx(rates.median())
        assert c["market_p75"] == pytest.approx(rates.quantile(0.75))
        assert c["market_p25"] == pytest.approx(rates.quantile(0.25))
        assert c["sample_size"] == len(rates)


def test_validated_metrics_and_flags(result):
    v = result.validated_df
    for _, r in v.iterrows():
        var = (r["prospect_rate"] - r["market_p50"]) / r["market_p50"] * 100
        assert r["raw_variance_pct"] == pytest.approx(var)
        assert r["variance_flag"] == variance_flag(var)
        assert r["annual_uplift_p50"] == pytest.approx(max(0.0, r["market_p50"] - r["prospect_rate"]) * r["annual_volume"])
        assert r["annual_uplift_p75"] == pytest.approx(max(0.0, r["market_p75"] - r["prospect_rate"]) * r["annual_volume"])
        assert r["prospect_pct_medicare"] == pytest.approx(r["prospect_rate"] / r["medicare_rate"] * 100)
        assert 0 <= r["market_percentile_rank"] <= 100
        assert re.fullmatch(r"MRF-\d+", r["sample_mrf_citation"])
    under = v["variance_flag"].isin(["CRITICAL_UNDER_MARKET", "MODERATE_UNDER_MARKET"]).mean()
    assert 0.5 <= under <= 0.7, "≈60% of matched lines should be under market"
    assert (v["variance_flag"] == "COMPETITIVE_BAND").any() and (v["variance_flag"] == "ABOVE_MARKET").any()


def test_citations_point_at_real_records(engine, result):
    mrf = engine.mrf_df.set_index("mrf_record_id")
    for _, r in result.validated_df.iterrows():
        rec = mrf.loc[r["sample_mrf_citation"]]
        assert rec["cpt_code"] == r["cpt_code"] and rec["payer_name"] == r["payer_name"]
        assert rec["hospital_name"] == r["citation_hospital"]


def test_dnf_resolution(engine, result):
    d = result.dnf_df
    for _, r in d.iterrows():
        reason, proxy, factor = EXPECTED_DNF[r["cpt_code"]]
        assert r["dnf_reason_code"] == reason
        assert r["proxy_cpt"] == proxy
        assert r["adjustment_factor"] == pytest.approx(factor)
        cohort = engine.cohort_df[(engine.cohort_df["cpt_code"] == proxy) & (engine.cohort_df["payer_name"] == r["payer_name"])]
        assert r["estimated_market_rate"] == pytest.approx(cohort["market_p50"].iloc[0] * factor)
        assert r["estimated_market_rate"] > 0
        assert r["annual_dnf_recovery"] == pytest.approx(max(0.0, r["estimated_market_rate"] - r["prospect_rate"]) * r["annual_volume"])
        assert r["annual_dnf_recovery"] >= 0
        assert r["confidence_tier"] and r["resolution_methodology"] and r["proxy_mrf_citation"].startswith("MRF-")
    assert d[d["cpt_code"] == "70551"]["modifier"].eq("26").all()


def test_dnf_factor_override_changes_recovery(engine, result):
    base = result.dnf_df[[c for c in PROSPECT_COLUMNS]]
    rid = result.dnf_df.iloc[0]["prospect_row_id"]
    lower = engine.resolve_dnf_queue(base, {rid: 0.5})
    assert lower.iloc[0]["adjustment_factor"] == 0.5
    assert lower.iloc[0]["annual_dnf_recovery"] < result.dnf_df.iloc[0]["annual_dnf_recovery"]


def test_cfo_story_is_compelling(kpis):
    assert kpis["validated_uplift_p50"] > 1_000_000
    assert 1_000_000 <= kpis["total_addressable"] <= 10_000_000
    assert kpis["realized_uplift"] == pytest.approx(kpis["total_addressable"] * 0.5)
    assert kpis["payback_months"] == pytest.approx(SUB / kpis["realized_uplift"] * 12)
    assert kpis["match_rate_pct"] == pytest.approx(52 / 60 * 100)
    assert kpis["coverage_with_dnf_pct"] == pytest.approx(100.0)
    assert kpis["prospect_pct_medicare"] < kpis["market_pct_medicare"]


def test_scenarios_and_sensitivity(result):
    v, d = result.validated_df, result.dnf_df
    p50 = compute_kpis(v, d)
    p75 = compute_kpis(v, d, TARGET_P75)
    floor = compute_kpis(v, d, TARGET_FLOOR, floor_pct=195)
    no_dnf = compute_kpis(v, d, include_dnf_recovery=False)
    assert p75["total_addressable"] > p50["total_addressable"]
    assert floor["validated_scenario_gap"] == pytest.approx(
        ((v["medicare_rate"] * 1.95 - v["prospect_rate"]).clip(lower=0) * v["annual_volume"]).sum())
    assert no_dnf["total_addressable"] == pytest.approx(p50["validated_uplift_p50"])
    sm = sensitivity_matrix(v, d, SUB)
    assert len(sm) == 12
    par = pareto_table(v, v["annual_uplift_p50"])
    assert par["cumulative_share_pct"].iloc[-1] == pytest.approx(100.0)


def test_custom_upload_with_unknown_code_is_not_dropped(engine):
    df = engine.default_prospect().head(6).copy()
    extra = df.iloc[[0]].copy()
    extra["prospect_row_id"], extra["cpt_code"] = "X-1", "99999"
    res = engine.run_validation(pd.concat([df, extra]))
    assert len(res.validated_df) + len(res.dnf_df) == 7
    unknown = res.dnf_df[res.dnf_df["cpt_code"] == "99999"].iloc[0]
    assert unknown["dnf_reason_code"] == "NO_CROSSWALK_DEFINED"
    assert unknown["annual_dnf_recovery"] == 0


def test_bad_upload_rejected():
    with pytest.raises(ValueError, match="missing required columns"):
        load_prospect_csv(io.StringIO("cpt_code,prospect_rate\n27447,100\n"))


# --------------------------------------------------------------------------- 3. Excel
def test_workbook_sheets_and_live_formulas(xlsx):
    wb = load_workbook(io.BytesIO(xlsx), data_only=False)
    assert wb.sheetnames == [SHEET_EXEC, SHEET_BENCH, SHEET_DNF]
    ex, be, dn = wb[SHEET_EXEC], wb[SHEET_BENCH], wb[SHEET_DNF]
    for cell in ("C12", "C13", "C14", "C15", "C16", "C17", "C18", "C19", "C20", "C21"):
        assert str(ex[cell].value).startswith("="), cell
    for cell in ("C6", "C7", "C8", "C9"):
        assert not str(ex[cell].value).startswith("="), f"{cell} must be an editable input"
    for col in ("J", "K", "L", "O", "P", "Q"):
        for row in range(2, 54):
            assert str(be[f"{col}{row}"].value).startswith("="), f"{col}{row}"
    for cell in ("J54", "K54", "L54"):
        assert str(be[cell].value).startswith("=SUM("), cell
    assert ex["C13"].value == f"=SUM({SHEET_BENCH}!K2:K53)"
    for row in range(2, 10):
        for col in ("F", "K", "M"):
            assert str(dn[f"{col}{row}"].value).startswith("="), f"{col}{row}"
    assert dn["K10"].value == "=SUM(K2:K9)"
    assert count_formulas(xlsx) > 300


def test_workbook_math_equals_engine(xlsx, kpis, result):
    ev = WorkbookEvaluator(xlsx)
    assert ev.value(SHEET_EXEC, "C12") == pytest.approx(kpis["current_contracted_revenue"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C13") == pytest.approx(kpis["validated_uplift_p50"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C14") == pytest.approx(kpis["validated_uplift_p75"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C15") == pytest.approx(kpis["dnf_recovery"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C16") == pytest.approx(kpis["total_addressable"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C17") == pytest.approx(kpis["realized_uplift"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C18") == pytest.approx(kpis["year1_net_benefit"], abs=0.01)
    assert ev.value(SHEET_EXEC, "C19") == pytest.approx(kpis["roi_multiple"])
    assert ev.value(SHEET_EXEC, "C20") == pytest.approx(kpis["payback_months"])
    assert ev.value(SHEET_EXEC, "C21") == pytest.approx(kpis["three_year_net_value"], abs=0.01)
    assert ev.value(SHEET_DNF, "K10") == pytest.approx(result.dnf_df["annual_dnf_recovery"].sum(), abs=0.01)


def test_workbook_recalculates_when_cfo_edits_assumption(xlsx):
    wb = load_workbook(io.BytesIO(xlsx))
    wb[SHEET_EXEC]["C7"] = 0.25
    wb[SHEET_EXEC]["C8"] = 0
    buf = io.BytesIO()
    wb.save(buf)
    ev = WorkbookEvaluator(buf.getvalue())
    assert ev.value(SHEET_EXEC, "C15") == 0
    assert ev.value(SHEET_EXEC, "C17") == pytest.approx(ev.value(SHEET_EXEC, "C13") * 0.25)


# --------------------------------------------------------------------------- 4. agent + skill
def test_skill_file_parses():
    meta, body = load_skill()
    assert meta["name"] == "trek-benchmark-validation-analyst"
    assert meta["version"] == "2.1.0"
    assert meta["citation_policy"] == "strict_mrf_record_id"
    assert set(meta["tools"]) == {"duckdb_mrf_benchmark_lookup", "dnf_crosswalk_resolver", "openpayer_policy_checker",
                                  "cfo_roi_memo_drafter"}
    for rule in ("Rule 1", "Rule 2", "Rule 3", "Rule 4"):
        assert rule in body


@pytest.mark.parametrize("label", list(PRESET_PROMPTS))
def test_agent_presets_are_cited_and_two_altitude(label, engine, result, kpis):
    text, info = ask_benchmark_agent(PRESET_PROMPTS[label], result.validated_df, result.dnf_df, kpis,
                                     engine.policies, api_key="")
    assert info["mode_used"] == "deterministic"
    assert info["skill_name"] == "trek-benchmark-validation-analyst"
    assert TECH_HEADER in text and CFO_HEADER in text
    assert text.index(TECH_HEADER) < text.index(CFO_HEADER)
    assert "[MRF-" in text and info["citations_count"] > 0
    assert re.search(r"\b\d{4}[0-9T]\b", text)
    cited = set(re.findall(r"MRF-\d+", text))
    assert cited <= set(engine.mrf_df["mrf_record_id"]), "agent cited a record that does not exist"


def test_agent_preset_content(engine, result, kpis):
    v, d, pol = result.validated_df, result.dnf_df, engine.policies
    ask = lambda q: ask_benchmark_agent(q, v, d, kpis, pol, api_key="")[0]  # noqa: E731
    top = ask(PRESET_PROMPTS["🔥 Top 5 Renegotiation Targets"])
    expected_top = v.groupby("cpt_code")["annual_uplift_p50"].sum().nlargest(5).index
    for cpt in expected_top:
        assert cpt in top
    dnf = ask(PRESET_PROMPTS["🧩 Defend DNF Methodology"])
    for cpt in EXPECTED_DNF:
        assert cpt in dnf
    for reason, _, _ in EXPECTED_DNF.values():
        assert reason in dnf
    policy = ask(PRESET_PROMPTS["🛡️ OpenPayer Policy + Rate Risk"])
    assert "27447" in policy and "CO-197" in policy and "Double exposure" in policy
    memo = ask(PRESET_PROMPTS["📄 Draft CFO Board Memo"])
    assert "MEMO" in memo and ("70553" in memo or "27447" in memo)


@pytest.mark.parametrize("query,intent,needle", [
    ("Tell me about 27447", "cpt_lookup", "27447"),
    ("Compare UHC vs Aetna in Radiology", "payer_compare", "Radiology"),
    ("Compare our Orthopedics and Radiology rates as a % of Medicare", "service_line_compare", "Orthopedics"),
    ("Explain 0408T", "cpt_lookup", "EMERGING_CAT_III_NO_MRF_RATE"),
    ("hello", "overview", "[MRF-"),
])
def test_agent_freeform_routing(query, intent, needle, engine, result, kpis):
    text, info = ask_benchmark_agent(query, result.validated_df, result.dnf_df, kpis, engine.policies, api_key="")
    assert info["intent"] == intent
    assert needle in text
    assert TECH_HEADER in text and CFO_HEADER in text


def test_agent_live_failure_falls_back(monkeypatch, engine, result, kpis):
    def boom(*a, **k):
        raise ConnectionError("offline")
    monkeypatch.setattr(agent_mod, "_ask_live", boom)
    text, info = ask_benchmark_agent("Tell me about 27447", result.validated_df, result.dnf_df, kpis, engine.policies,
                                     api_key="sk-test")
    assert info["mode_used"] == "deterministic" and "offline" in info["fallback_reason"]
    assert "[MRF-" in text


def test_agent_handles_empty_scope(engine, result, kpis):
    empty_v = result.validated_df.iloc[0:0]
    empty_d = result.dnf_df.iloc[0:0]
    for q in PRESET_PROMPTS.values():
        text, _ = ask_benchmark_agent(q, empty_v, empty_d, kpis, engine.policies, api_key="")
        assert TECH_HEADER in text and CFO_HEADER in text


# --------------------------------------------------------------------------- 5. demo script stays in sync
def test_demo_script_matches_live_numbers():
    facts = default_facts()
    assert SCRIPT_PATH.exists(), "run: python -m src.demo_script"
    assert SCRIPT_PATH.read_text(encoding="utf-8") == render_markdown(facts), \
        "DEMO_SCRIPT.md is stale — run: python -m src.demo_script"
    assert "{" not in render_markdown(facts).split("## Part 2")[1]
    assert all(v != "—" for v in facts.values()), "a talk-track fact could not be computed"
    assert len(SLIDES) == 9


# --------------------------------------------------------------------------- 6. UI (headless)
def test_app_renders_every_slide_and_workbench():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120).run()
    assert not at.exception, at.exception
    for s in SLIDES:
        at.session_state["slide_nav"] = s["nav"]
        at.run()
        assert not at.exception, (s["nav"], at.exception)
        assert any(s["title"].split(":")[0][:20] in m.value for m in at.markdown), s["nav"]
    at.session_state["slide_nav"] = "8 · Agent"
    at.run()
    at.button(key="preset_2").click().run()
    assert not at.exception
    assert any("Double exposure" in m.value for m in at.markdown)
    at.sidebar.radio(key="mode").set_value("🧰 Workbench").run()
    assert not at.exception
    assert len(at.tabs) == 5


def test_app_levers_drive_kpis():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120).run()
    at.session_state["slide_nav"] = "6 · ROI"
    at.run()
    before = [m.value for m in at.markdown if "Realized uplift" in m.value][0]
    at.slider(key="in_real").set_value(0.25).run()
    assert not at.exception
    after = [m.value for m in at.markdown if "Realized uplift" in m.value][0]
    assert "@ 25%" in after and before != after


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v", "-p", "no:cacheprovider"]))
